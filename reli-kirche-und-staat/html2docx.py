#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
html2docx.py — Kopiervorlagen-HTML in ein echtes Word-Dokument (.docx) umwandeln.

Aufruf:
    python3 html2docx.py lessons/0002-ds1-kopiervorlagen.html word/DS1-Kopiervorlagen.docx

Das Skript ist auf die Struktur der Kopiervorlagen-Dateien dieses Projekts
zugeschnitten (identische CSS-Klassen in allen Doppelstunden 1-6):

    <section>            -> eigene Seite (Seitenumbruch davor)
      .sec-label         -> kleine Kennzeile ueber der Ueberschrift
                            (enthaelt sie "nicht kopieren", wird sie zum Warnkasten)
      h2                 -> Ueberschrift 1   (Navigationsbereich in Word)
      h3                 -> Ueberschrift 2
      h4                 -> Ueberschrift 3
      .blatt             -> Arbeitsblatt-Container
      .blatt > header    -> Kopfzeile mit .id / .who
      .folie             -> Projektionsfolie (mit .folie-nr, Faksimile oder Bild)
      .auftragsbox       -> Auftrag, <ol> wird zur nummerierten Liste
      .rechercheblock    -> Rechercheauftrag (nummerierte Liste + Quellenzeile)
      .hinweis/.warnung  -> Kasten mit Rahmen und Hintergrundfarbe
      .namensfeld        -> "Gruppe ____  Namen ____  Datum ____"
      table.raster       -> beschreibbares Raster (hohe Zellen, Rahmen)
      table.verlauf      -> Verlaufstabelle (Rahmen)
      table.fax-*        -> Faksimile-Tabellen (Rahmen, kleinere Schrift)
      .streitsatz        -> nummerierter Streitsatz
      .schild            -> Eckenschild, eigene Seite, sehr grosse Schrift
      footer             -> Bildnachweis; Links werden mit voller URL ausgegeben

Nur die Standardbibliothek plus python-docx wird benutzt.
Bilder duerfen als data:-URI eingebettet oder als Dateipfad verlinkt sein.
"""

import base64
import binascii
import html.parser
import io
import os
import re
import sys
import urllib.parse

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

# --------------------------------------------------------------------------
# 1. Minimaler HTML-Parser -> Baum aus Element-/Textknoten
# --------------------------------------------------------------------------

VOID_TAGS = {"br", "img", "hr", "meta", "link", "input", "col", "source"}
# Diese Tags samt Inhalt werden komplett verworfen.
DROP_TAGS = {"script", "style", "head", "title", "nav"}


class Element:
    __slots__ = ("tag", "attrs", "children", "parent")

    def __init__(self, tag, attrs=None, parent=None):
        self.tag = tag
        self.attrs = attrs or {}
        self.children = []
        self.parent = parent

    # -- kleine Helfer -----------------------------------------------------
    @property
    def cls(self):
        return set(self.attrs.get("class", "").split())

    def has(self, name):
        return name in self.cls

    def find(self, tag=None, cls=None):
        """Erster Nachfahre mit passendem Tag und/oder Klasse."""
        for node in self.iter_elements():
            if tag and node.tag != tag:
                continue
            if cls and cls not in node.cls:
                continue
            return node
        return None

    def iter_elements(self):
        for child in self.children:
            if isinstance(child, Element):
                yield child
                yield from child.iter_elements()

    def text(self):
        out = []
        for child in self.children:
            if isinstance(child, str):
                out.append(child)
            elif child.tag == "br":
                out.append(" ")
            else:
                out.append(child.text())
        return "".join(out)

    def __repr__(self):  # pragma: no cover - nur zum Debuggen
        return "<%s class=%r>" % (self.tag, self.attrs.get("class", ""))


class TreeBuilder(html.parser.HTMLParser):
    """Baut aus dem HTML einen Elementbaum. convert_charrefs loest &shy; etc. auf."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Element("root")
        self.stack = [self.root]
        self.dropping = 0  # Tiefenzaehler innerhalb von DROP_TAGS

    def handle_starttag(self, tag, attrs):
        if self.dropping:
            if tag in DROP_TAGS:
                self.dropping += 1
            return
        if tag in DROP_TAGS:
            self.dropping = 1
            return
        node = Element(tag, dict(attrs), self.stack[-1])
        self.stack[-1].children.append(node)
        if tag not in VOID_TAGS:
            self.stack.append(node)

    def handle_startendtag(self, tag, attrs):
        if self.dropping:
            return
        node = Element(tag, dict(attrs), self.stack[-1])
        self.stack[-1].children.append(node)

    def handle_endtag(self, tag):
        if self.dropping:
            if tag in DROP_TAGS:
                self.dropping -= 1
            return
        if tag in VOID_TAGS:
            return
        # bis zum passenden offenen Element zurueckwandern (tolerant gegen
        # fehlende Endtags, wie sie in handgeschriebenem HTML vorkommen)
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                return

    def handle_data(self, data):
        if self.dropping:
            return
        self.stack[-1].children.append(data)


def parse_html(path):
    with open(path, "r", encoding="utf-8") as fh:
        source = fh.read()
    parser = TreeBuilder()
    parser.feed(source)
    parser.close()
    return parser.root


# --------------------------------------------------------------------------
# 2. Text-Normalisierung
# --------------------------------------------------------------------------

SOFT_HYPHEN = "­"


def clean(text):
    """HTML-Whitespace zusammenfassen, weiche Trennstriche entfernen.

    Weiche Trennstriche (&shy;) stehen im HTML nur fuer den Zeilenumbruch in
    Tabellenkoepfen; in Word wuerden sie beim Kopieren stoeren.
    """
    text = text.replace(SOFT_HYPHEN, "")
    text = re.sub(r"[ \t\r\n]+", " ", text)
    return text


# --------------------------------------------------------------------------
# 3. Low-Level-Helfer fuer docx-XML, das python-docx nicht direkt anbietet
# --------------------------------------------------------------------------

ACCENT = RGBColor(0x1E, 0x3A, 0x5F)   # dunkles Blau  (wie --accent im HTML)
INK_SOFT = RGBColor(0x4B, 0x55, 0x63)
INK_FAINT = RGBColor(0x6B, 0x72, 0x80)
FLAG = RGBColor(0x8C, 0x2A, 0x24)     # Rot fuer Warnungen
EPOCH = RGBColor(0x6A, 0x59, 0x36)    # Braun fuer Hinweise
NOW = RGBColor(0x1F, 0x5B, 0x4E)      # Gruen fuer Recherche

FILL_ACCENT = "E5ECF4"
FILL_EPOCH = "F0EBDF"
FILL_NOW = "E2EDEA"
FILL_FLAG = "F6E7E5"
FILL_GREY = "EFF1F4"


def shade(element, fill):
    """Hintergrundfarbe fuer eine Zelle (w:tcPr) oder Zeile setzen."""
    pr = element.get_or_add_tcPr() if hasattr(element, "get_or_add_tcPr") else element
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def shade_cell(cell, fill):
    shade(cell._tc, fill)


def shade_paragraph(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def paragraph_border(paragraph, edges=("left",), color="1E3A5F", size=18, space=6):
    """Farbige Linie an einer Absatzkante — ersetzt die CSS-Rahmen der Kaesten."""
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    for edge in edges:
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), str(space))
        el.set(qn("w:color"), color)
        pbdr.append(el)
    pPr.append(pbdr)


def cell_borders(cell, color="9AA3AE", size=6):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def set_row_height(row, cm, exact=False):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(cm * 567)))  # 1 cm = 567 Twips
    h.set(qn("w:hRule"), "exact" if exact else "atLeast")
    trPr.append(h)


def keep_with_next(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:keepNext"))


def style_num_id(doc, style_id):
    """numId, den eine Listen-Formatvorlage benutzt (z. B. ListNumber)."""
    for style in doc.styles.element.findall(qn("w:style")):
        if style.get(qn("w:styleId")) != style_id:
            continue
        num_id = style.find(qn("w:pPr") + "/" + qn("w:numPr") + "/" + qn("w:numId"))
        if num_id is not None:
            return int(num_id.get(qn("w:val")))
    return None


def restart_numbering(doc, style_id="ListNumber"):
    """Neue Nummerierungs-Instanz anlegen, damit jede Liste wieder bei 1 beginnt.

    Alle Absaetze mit der Vorlage "Liste mit Nummern" teilen sich in der
    Word-Standardvorlage denselben numId — ohne diesen Schritt wuerde der
    zweite Arbeitsauftrag mit 5. weiterzaehlen statt mit 1.
    Rueckgabe: der neue numId oder None, wenn die Vorlage nicht gefunden wurde.
    """
    try:
        numbering = doc.part.numbering_part.element
    except (NotImplementedError, KeyError, AttributeError):
        return None
    base = style_num_id(doc, style_id)
    if base is None:
        return None
    abstract_id, max_id = None, 0
    for num in numbering.findall(qn("w:num")):
        this_id = int(num.get(qn("w:numId")))
        max_id = max(max_id, this_id)
        if this_id == base:
            ref = num.find(qn("w:abstractNumId"))
            if ref is not None:
                abstract_id = ref.get(qn("w:val"))
    if abstract_id is None:
        return None
    new = OxmlElement("w:num")
    new.set(qn("w:numId"), str(max_id + 1))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), abstract_id)
    new.append(ref)
    numbering.append(new)
    return max_id + 1


def apply_num_id(paragraph, num_id, level=0):
    pPr = paragraph._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    numPr.append(ilvl)
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    numPr.append(nid)
    pPr.append(numPr)


def add_hyperlink(paragraph, url, text, size=None):
    """Echter Word-Hyperlink (python-docx hat dafuer keine API)."""
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1E3A5F")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        rPr.append(sz)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


# --------------------------------------------------------------------------
# 4. Bilder
# --------------------------------------------------------------------------

DATA_URI = re.compile(r"^data:image/(?P<ext>[a-zA-Z0-9.+-]+);base64,(?P<payload>.*)$", re.S)

# Minimales JFIF-APP0-Segment (16 Byte Nutzlast): Kennung, Version 1.1,
# Einheit 0, x/y-Dichte 1, keine Thumbnail-Daten.
JFIF_APP0 = (b"\xff\xe0\x00\x10JFIF\x00\x01\x01\x00\x00\x01\x00\x01\x00\x00")


def normalize_jpeg(blob):
    """Bare JPEGs um ein JFIF-Segment ergaenzen.

    Die Fotos in den Vorlagen sind mit Optimierern erzeugt worden und beginnen
    direkt mit dem Quantisierungstabellen-Marker (FFD8 FFDB). python-docx
    erkennt nur JPEGs mit APP0-(JFIF-) oder APP1-(Exif-)Segment und wirft sonst
    UnrecognizedImageError. Das Segment wird deshalb nachtraeglich eingesetzt;
    die Bilddaten selbst bleiben unveraendert.
    """
    if len(blob) < 4 or blob[:2] != b"\xff\xd8":
        return blob
    if blob[2:4] in (b"\xff\xe0", b"\xff\xe1"):
        return blob
    return blob[:2] + JFIF_APP0 + blob[2:]


def image_stream(src, base_dir):
    """Liefert (BytesIO, Kurzname) fuer ein <img src> — data:-URI oder Datei."""
    src = src.strip()
    match = DATA_URI.match(src)
    if match:
        payload = re.sub(r"\s+", "", match.group("payload"))
        try:
            return io.BytesIO(normalize_jpeg(base64.b64decode(payload))), "eingebettet"
        except (binascii.Error, ValueError):
            return None, None
    if src.startswith(("http://", "https://")):
        return None, None  # kein Netzzugriff: extern verlinkte Bilder werden uebersprungen
    path = os.path.normpath(os.path.join(base_dir, urllib.parse.unquote(src)))
    if os.path.isfile(path):
        with open(path, "rb") as fh:
            return io.BytesIO(normalize_jpeg(fh.read())), os.path.basename(path)
    return None, None


# --------------------------------------------------------------------------
# 5. Renderer
# --------------------------------------------------------------------------

INLINE_TAGS = {"a", "span", "strong", "b", "em", "i", "code", "small", "u",
               "sub", "sup", "br", "abbr", "cite", "mark", "time", "s"}
# span-Klassen, die im HTML zwar inline stehen, aber wie ein eigener Absatz wirken
BLOCK_SPAN_CLASSES = {"quelle", "t", "folie-nr"}

BODY_FONT = "Calibri"
MONO_FONT = "Consolas"
DISPLAY_FONT = "Cambria"


class Converter:
    def __init__(self, doc, base_dir):
        self.doc = doc
        self.base_dir = base_dir
        self.show_urls = False   # im Bildnachweis werden URLs mitgedruckt
        self.first_section = True
        self.images = 0
        self.skipped_images = 0

    # ---- Absatz-Grundlagen ----------------------------------------------
    def para(self, style=None, space_before=None, space_after=6, align=None):
        p = self.doc.add_paragraph(style=style)
        fmt = p.paragraph_format
        if space_before is not None:
            fmt.space_before = Pt(space_before)
        if space_after is not None:
            fmt.space_after = Pt(space_after)
        if align is not None:
            fmt.alignment = align
        return p

    def label(self, text, color=INK_FAINT, size=8.5, bold=True, space_after=2):
        p = self.para(space_after=space_after)
        run = p.add_run(clean(text).strip().upper())
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = MONO_FONT
        return p

    # ---- Inline-Formatierung --------------------------------------------
    def inline(self, paragraph, node, fmt):
        """Fuegt den Inhalt eines Inline-Knotens als Runs in den Absatz ein."""
        for child in node.children:
            if isinstance(child, str):
                text = clean(child)
                if text:
                    self.add_run(paragraph, text, fmt)
                continue
            tag, cls = child.tag, child.cls
            if tag == "br":
                if paragraph.runs:
                    paragraph.runs[-1].add_break(WD_BREAK.LINE)
                else:
                    self.add_run(paragraph, "", fmt).add_break(WD_BREAK.LINE)
                continue
            sub = dict(fmt)
            if tag in ("strong", "b"):
                sub["bold"] = True
            elif tag in ("em", "i", "cite"):
                sub["italic"] = True
            elif tag == "code":
                sub["mono"] = True
            elif tag == "u":
                sub["underline"] = True
            elif tag == "small":
                sub["size"] = (sub.get("size") or 11) - 1.5
            if "hl" in cls:  # em.hl = im HTML farbig hervorgehoben
                sub["bold"] = True
                sub["color"] = FLAG
                sub["italic"] = False
            if tag == "a":
                href = child.attrs.get("href", "")
                text = clean(child.text()).strip()
                if href.startswith(("http://", "https://")):
                    add_hyperlink(paragraph, href, text, size=sub.get("size"))
                    if self.show_urls:
                        url_fmt = dict(sub, mono=True, size=8,
                                       color=INK_FAINT, bold=False, italic=False)
                        self.add_run(paragraph, " (%s)" % href, url_fmt)
                else:
                    self.add_run(paragraph, text, sub)
                continue
            self.inline(paragraph, child, sub)

    def add_run(self, paragraph, text, fmt):
        run = paragraph.add_run(text)
        font = run.font
        font.bold = bool(fmt.get("bold"))
        font.italic = bool(fmt.get("italic"))
        if fmt.get("underline"):
            font.underline = True
        if fmt.get("mono"):
            font.name = MONO_FONT
        elif fmt.get("font"):
            font.name = fmt["font"]
        if fmt.get("size"):
            font.size = Pt(fmt["size"])
        if fmt.get("color") is not None:
            font.color.rgb = fmt["color"]
        return run

    def is_inline(self, child):
        if isinstance(child, str):
            return True
        if child.tag not in INLINE_TAGS:
            return False
        if child.tag == "span" and (child.cls & BLOCK_SPAN_CLASSES):
            return False
        return True

    # ---- Block-Ebene -----------------------------------------------------
    def children(self, node, fmt=None, style=None, container=None):
        """Kinder eines Block-Containers rendern.

        Aufeinanderfolgende Inline-Knoten werden zu je einem Absatz gebuendelt,
        Block-Knoten einzeln weitergereicht. So funktioniert auch gemischter
        Inhalt (Text + <span> + <div>) wie in .fax-gesetz.
        """
        fmt = fmt or {}
        buffer = []

        def flush():
            if not buffer:
                return
            if not clean("".join(
                    c if isinstance(c, str) else c.text() for c in buffer)).strip():
                buffer.clear()
                return
            wrapper = Element("p")
            wrapper.children = list(buffer)
            buffer.clear()
            p = self.paragraph_in(container, style)
            self.inline(p, wrapper, fmt)

        for child in node.children:
            if self.is_inline(child):
                buffer.append(child)
            else:
                flush()
                self.block(child, fmt=fmt, container=container)
        flush()

    def paragraph_in(self, container, style=None):
        if container is None:
            return self.para(style=style)
        p = container.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(4)
        return p

    def block(self, node, fmt=None, container=None):
        """Ein Block-Element rendern; Verzweigung nach CSS-Klasse, dann Tag."""
        fmt = fmt or {}
        cls = node.cls
        tag = node.tag

        # --- Sonderfaelle nach Klasse ------------------------------------
        if "schild" in cls:
            return self.render_schild(node)
        if "streitsatz" in cls:
            return self.render_streitsatz(node)
        if "namensfeld" in cls:
            return self.render_namensfeld(node)
        if "auftragsbox" in cls:
            return self.render_auftragsbox(node)
        if "rechercheblock" in cls:
            return self.render_kasten(node, FILL_NOW, NOW, "9CC2B7")
        if "hinweis" in cls:
            return self.render_kasten(node, FILL_EPOCH, EPOCH, "C9BC98")
        if "warnung" in cls:
            return self.render_kasten(node, FILL_FLAG, FLAG, "D8A6A1")
        if "sec-label" in cls:
            return self.render_sec_label(node)
        if "folie" in cls:
            return self.render_folie(node)
        if "blatt" in cls:
            return self.render_blatt(node)
        if "folien-was" in cls:
            p = self.para(space_after=12)
            self.inline(p, node, dict(fmt, size=9.5, color=INK_SOFT))
            return None
        if "bildquelle" in cls or "quellen" in cls:
            p = self.para(space_after=6)
            self.inline(p, node, dict(fmt, size=9, color=INK_SOFT, mono=True))
            return None
        if "fax-gesetz" in cls:
            return self.children(node, fmt=dict(fmt, size=13, font=DISPLAY_FONT),
                                 container=container)
        if "fax-para" in cls:
            return self.children(node, fmt=dict(fmt, size=9.5, mono=True,
                                                color=INK_SOFT), container=container)
        if node.tag == "span" and "quelle" in cls:
            p = self.paragraph_in(container)
            self.inline(p, node, dict(fmt, size=8.5, mono=True, color=INK_FAINT))
            return None
        if node.tag == "span" and "folie-nr" in cls:
            return self.label("Folie " + clean(node.text()).strip(), color=ACCENT)

        # --- Sonderfaelle nach Tag ---------------------------------------
        if tag == "img":
            return self.render_img(node, container)
        if tag == "table":
            return self.render_table(node, container)
        if tag in ("ul", "ol"):
            return self.render_list(node, ordered=(tag == "ol"), fmt=fmt,
                                    container=container)
        if tag == "li":  # kommt nur bei kaputtem Markup vor
            return self.children(node, fmt=fmt, container=container)
        if tag == "p":
            p = self.paragraph_in(container)
            self.inline(p, node, fmt)
            return None
        if tag in ("h1", "h2", "h3", "h4", "h5"):
            return self.render_heading(node, tag)
        if tag == "hr":
            return None
        if tag == "section":
            return self.render_section(node)
        if tag == "footer":
            prev = self.show_urls
            self.show_urls = True       # Lizenz-URLs im Bildnachweis mitdrucken
            self.children(node, fmt=fmt, container=container)
            self.show_urls = prev
            return None
        if tag == "header" and node.parent is not None and node.parent.has("blatt"):
            return self.render_blatt_header(node)

        # --- generischer Container ---------------------------------------
        return self.children(node, fmt=fmt, container=container)

    # ---- konkrete Renderer ----------------------------------------------
    def render_section(self, node):
        if not self.first_section:
            self.doc.add_page_break()
        self.first_section = False
        self.children(node)

    def render_heading(self, node, tag):
        # h1 = Titel, h2 = Ueberschrift 1, h3 = Ueberschrift 2, h4 = Ueberschrift 3
        level = {"h1": 0, "h2": 1, "h3": 2, "h4": 3, "h5": 4}[tag]
        # Die Bildnachweis-Ueberschrift steht im HTML als h3 im <footer>,
        # ist aber ein eigener Hauptabschnitt -> Ueberschrift 1.
        if tag == "h3" and node.parent is not None and node.parent.tag == "footer":
            level = 1
        p = self.doc.add_heading("", level=level)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.space_before = Pt(10 if level else 0)
        self.inline(p, node, {})
        keep_with_next(p)
        return None

    def render_sec_label(self, node):
        text = clean(node.text()).strip()
        if "nicht kopieren" in text.lower():
            p = self.para(space_before=4, space_after=8)
            run = p.add_run(text.upper())
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = FLAG
            run.font.name = MONO_FONT
            shade_paragraph(p, FILL_FLAG)
            paragraph_border(p, edges=("top", "bottom", "left", "right"),
                             color="8C2A24", size=8)
            keep_with_next(p)
            return None
        p = self.label(text, color=ACCENT, size=9)
        keep_with_next(p)
        return None

    def render_blatt_header(self, node):
        ident = node.find(cls="id")
        who = node.find(cls="who")
        p = self.para(space_before=8, space_after=8)
        if ident is not None:
            run = p.add_run(clean(ident.text()).strip())
            run.font.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = ACCENT
            run.font.name = MONO_FONT
        if who is not None:
            run = p.add_run("   " + clean(who.text()).strip())
            run.font.size = Pt(9)
            run.font.color.rgb = INK_FAINT
            run.font.name = MONO_FONT
        paragraph_border(p, edges=("bottom",), color="1E3A5F", size=12)
        keep_with_next(p)
        return None

    def render_blatt(self, node):
        # Das .blatt ist im Druck die Seite selbst — kein zusaetzlicher Rahmen,
        # damit der Inhalt notfalls auf die naechste Seite umbrechen kann.
        self.children(node)

    def render_auftragsbox(self, node):
        title = node.find(cls="t")
        if title is not None:
            p = self.label(clean(title.text()).strip(), color=ACCENT, size=9)
            shade_paragraph(p, FILL_ACCENT)
            keep_with_next(p)
        for child in node.children:
            if isinstance(child, Element) and child.has("t"):
                continue
            if isinstance(child, str):
                continue
            self.block(child)
        return None

    def render_kasten(self, node, fill, color, border):
        """hinweis / warnung / rechercheblock als einzellige Tabelle mit Rahmen."""
        table = self.doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
        shade_cell(cell, fill)
        cell_borders(cell, color=border, size=8)
        title = node.find(cls="t")
        if title is not None:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(clean(title.text()).strip().upper())
            run.font.size = Pt(8.5)
            run.font.bold = True
            run.font.name = MONO_FONT
            run.font.color.rgb = color
        for child in node.children:
            if isinstance(child, Element) and child.has("t"):
                continue
            if isinstance(child, str):
                if not clean(child).strip():
                    continue
                p = cell.add_paragraph()
                p.add_run(clean(child).strip())
                continue
            self.block(child, container=cell)
        if not cell.paragraphs:
            cell.add_paragraph()
        self.para(space_after=6)  # Luft nach dem Kasten
        return None

    def render_namensfeld(self, node):
        parts = [clean(s.text()).strip()
                 for s in node.children
                 if isinstance(s, Element) and s.tag == "span"]
        if not parts:
            parts = [clean(node.text()).strip()]
        p = self.para(space_before=10, space_after=6)
        for i, part in enumerate(parts):
            if i:
                p.add_run("    ")
            run = p.add_run(part + ": ")
            run.font.size = Pt(9.5)
            run.font.color.rgb = INK_SOFT
            run.font.name = MONO_FONT
            line = p.add_run("_" * (26 if len(parts) < 3 else 20))
            line.font.size = Pt(9.5)
            line.font.color.rgb = INK_FAINT
        paragraph_border(p, edges=("top",), color="B7BEC7", size=6)
        return None

    def render_streitsatz(self, node):
        number = node.find(cls="n")
        satz = node.find(cls="s")
        p = self.para(space_before=6, space_after=8)
        p.paragraph_format.left_indent = Cm(0.4)
        if number is not None:
            run = p.add_run(clean(number.text()).strip() + "   ")
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = ACCENT
            run.font.name = DISPLAY_FONT
        if satz is not None:
            run = p.add_run(clean(satz.text()).strip())
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.name = DISPLAY_FONT
        paragraph_border(p, edges=("left",), color="1E3A5F", size=18, space=8)
        return None

    def render_schild(self, node):
        """Eckenschild: eigene Seite, zentriert, sehr grosse Schrift (Laminat)."""
        self.doc.add_page_break()
        klein = node.find(cls="klein")
        gross = node.find(cls="gross")
        for _ in range(3):
            self.para(space_after=0)
        if klein is not None:
            p = self.para(space_after=24, align=WD_ALIGN_PARAGRAPH.CENTER)
            run = p.add_run(clean(klein.text()).strip().upper())
            run.font.size = Pt(20)
            run.font.name = MONO_FONT
            run.font.color.rgb = INK_FAINT
        if gross is not None:
            p = self.para(space_after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.line_spacing = 1.0
            # <br> im HTML wird zum echten Zeilenumbruch
            lines, current = [], []
            for child in gross.children:
                if isinstance(child, str):
                    current.append(clean(child))
                elif child.tag == "br":
                    lines.append("".join(current))
                    current = []
                else:
                    current.append(clean(child.text()))
            lines.append("".join(current))
            for i, line in enumerate([l.strip() for l in lines if l.strip()]):
                run = p.add_run(line)
                run.font.size = Pt(72)          # deutlich ueber den geforderten 60 pt
                run.font.bold = True
                run.font.name = DISPLAY_FONT
                run.font.color.rgb = ACCENT
                if i < len([l for l in lines if l.strip()]) - 1:
                    run.add_break(WD_BREAK.LINE)
        return None

    def render_folie(self, node):
        nr = node.find(cls="folie-nr")
        if nr is not None:
            p = self.label("Folie " + clean(nr.text()).strip(), color=ACCENT, size=9)
            shade_paragraph(p, FILL_GREY)
            keep_with_next(p)
        for child in node.children:
            if isinstance(child, Element) and child.has("folie-nr"):
                continue
            if isinstance(child, str):
                if clean(child).strip():
                    p = self.para()
                    p.add_run(clean(child).strip())
                continue
            self.block(child)
        return None

    def render_img(self, node, container=None):
        src = node.attrs.get("src", "")
        stream, _name = image_stream(src, self.base_dir)
        alt = clean(node.attrs.get("alt", "")).strip()
        if stream is None:
            self.skipped_images += 1
            p = self.paragraph_in(container)
            run = p.add_run("[Bild konnte nicht eingebettet werden: %s]"
                            % (alt or src[:60]))
            run.font.italic = True
            run.font.color.rgb = FLAG
            return None
        if container is None:
            p = self.para(space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
        else:
            p = container.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(stream, width=Cm(14.5))
        self.images += 1
        if alt:
            cap = self.para(space_after=8, align=WD_ALIGN_PARAGRAPH.CENTER) \
                if container is None else container.add_paragraph()
            run = cap.add_run("Bildbeschreibung: " + alt)
            run.font.size = Pt(8.5)
            run.font.italic = True
            run.font.color.rgb = INK_FAINT
        return None

    def render_list(self, node, ordered, fmt, container=None):
        items = [c for c in node.children
                 if isinstance(c, Element) and c.tag == "li"]
        if not items:
            return None
        style = "List Number" if ordered else "List Bullet"
        # Jede nummerierte Liste bekommt eine eigene Nummerierung, sonst
        # zaehlt Word ueber das ganze Dokument durch.
        num_id = restart_numbering(self.doc) if (ordered and container is None) else None
        for index, item in enumerate(items, start=1):
            if container is None:
                p = self.doc.add_paragraph(style=style)
                p.paragraph_format.space_after = Pt(3)
                if num_id is not None:
                    apply_num_id(p, num_id)
            else:
                # In Tabellenzellen (Hinweis-/Recherchekaesten) sind
                # Listenvorlagen unzuverlaessig -> Marke von Hand setzen.
                p = container.add_paragraph()
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.left_indent = Cm(0.6)
                marker = ("%d. " % index) if ordered else "• "
                p.add_run(marker).font.bold = ordered
            self.inline(p, item, fmt)
        return None

    # ---- Tabellen --------------------------------------------------------
    def render_table(self, node, container=None):
        rows = []
        for section_tag in ("thead", "tbody", "tfoot"):
            for part in node.children:
                if isinstance(part, Element) and part.tag == section_tag:
                    rows.extend([(r, section_tag) for r in part.children
                                 if isinstance(r, Element) and r.tag == "tr"])
        rows.extend([(r, "tbody") for r in node.children
                     if isinstance(r, Element) and r.tag == "tr"])
        if not rows:
            return None

        caption = node.find(tag="caption")
        cls = node.cls
        is_raster = "raster" in cls
        is_fax = any(c.startswith("fax") for c in cls)

        # Spaltenzahl inkl. colspan bestimmen
        ncols = 0
        for row, _ in rows:
            width = sum(int(c.attrs.get("colspan", 1))
                        for c in row.children
                        if isinstance(c, Element) and c.tag in ("td", "th"))
            ncols = max(ncols, width)

        if caption is not None:
            p = self.label(clean(caption.text()).strip(), color=INK_FAINT, size=9)
            keep_with_next(p)

        if container is None:
            table = self.doc.add_table(rows=0, cols=ncols)
        else:
            table = container.add_table(rows=0, cols=ncols)
        table.style = "Table Grid"
        table.autofit = True

        for row, part in rows:
            cells = [c for c in row.children
                     if isinstance(c, Element) and c.tag in ("td", "th")]
            wrow = table.add_row()
            col = 0
            empty_row = True
            for cellnode in cells:
                span = int(cellnode.attrs.get("colspan", 1))
                if col >= ncols:
                    break
                target = wrow.cells[col]
                if span > 1 and col + span - 1 < ncols:
                    target = target.merge(wrow.cells[col + span - 1])
                text = clean(cellnode.text()).strip()
                if text and cellnode.tag == "td":
                    # Zeilenkoepfe (th) zaehlen nicht: eine Rasterzeile gilt als
                    # leer, wenn alle Datenzellen leer sind — dann wird sie hoch.
                    empty_row = False
                para = target.paragraphs[0]
                para.paragraph_format.space_after = Pt(2)
                cellfmt = {"size": 9 if is_fax else 10}
                header = (part == "thead") or cellnode.tag == "th"
                if header:
                    cellfmt["bold"] = True
                    cellfmt["color"] = ACCENT if part == "thead" else INK_SOFT
                if "z" in cellnode.cls:
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if "mark" in cellnode.cls or "mark" in row.cls:
                    cellfmt["bold"] = True
                    cellfmt["color"] = FLAG
                    shade_cell(target, FILL_FLAG)
                elif part == "thead":
                    shade_cell(target, FILL_ACCENT)
                elif cellnode.tag == "th":
                    shade_cell(target, FILL_GREY)
                self.inline(para, cellnode, cellfmt)
                target.vertical_alignment = WD_ALIGN_VERTICAL.TOP
                col += span

            if is_raster and part != "thead":
                # Beschreibbare Zeilen: leere Raster-Zeilen bekommen viel Platz,
                # Zeilen mit Ausfuellhinweis etwas weniger.
                set_row_height(wrow, 3.4 if empty_row else 2.4)

        # Spaltenbreiten: erste Spalte schmal, wenn sie Zeilenkoepfe enthaelt
        if is_raster and ncols > 1:
            usable = 17.0
            first = 2.8
            rest = (usable - first) / (ncols - 1)
            table.autofit = False
            for r in table.rows:
                cells = r.cells
                # Zeilen mit verbundenen Zellen (colspan) auslassen — dort
                # wuerde die Breitenzuweisung die Verbindung zerreissen.
                if len({id(c._tc) for c in cells}) != ncols:
                    continue
                for i, c in enumerate(cells):
                    c.width = Cm(first if i == 0 else rest)

        self.para(space_before=2, space_after=8)  # Abstand nach der Tabelle
        return None


# --------------------------------------------------------------------------
# 6. Dokument-Grundeinstellungen
# --------------------------------------------------------------------------

def prepare_document():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width = Cm(21.0)     # A4
    sec.page_height = Cm(29.7)
    for attr, value in (("top_margin", 1.8), ("bottom_margin", 1.8),
                        ("left_margin", 2.0), ("right_margin", 2.0)):
        setattr(sec, attr, Cm(value))

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    # Sprache auf Deutsch stellen, damit Word richtig trennt und prueft
    rpr = normal.element.get_or_add_rPr()
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "de-DE")
    rpr.append(lang)

    for name, size, color in (("Title", 22, ACCENT),
                              ("Heading 1", 17, ACCENT),
                              ("Heading 2", 13.5, None),
                              ("Heading 3", 11.5, None)):
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        style.font.name = DISPLAY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        if color is not None:
            style.font.color.rgb = color
    return doc


# --------------------------------------------------------------------------
# 7. Hauptprogramm
# --------------------------------------------------------------------------

def convert(src_path, out_path):
    root = parse_html(src_path)
    base_dir = os.path.dirname(os.path.abspath(src_path))
    doc = prepare_document()
    conv = Converter(doc, base_dir)

    sections = [n for n in root.iter_elements() if n.tag == "section"]
    masthead = next((n for n in root.iter_elements() if n.has("masthead")), None)

    # Titelblock aus dem <header class="masthead">
    if masthead is not None:
        kicker = masthead.find(cls="kicker")
        h1 = masthead.find(tag="h1")
        lede = masthead.find(cls="lede")
        if kicker is not None:
            conv.label(clean(kicker.text()).strip(), color=ACCENT, size=9)
        if h1 is not None:
            p = doc.add_heading("", level=0)
            conv.inline(p, h1, {})
        if lede is not None:
            p = conv.para(space_after=10)
            conv.inline(p, lede, {"size": 11, "color": INK_SOFT})
        conv.first_section = False  # naechste Section beginnt auf neuer Seite

    # Eckenschilder ans Ende: sie werden laminiert und einzeln aufgehaengt,
    # deshalb stehen sie hinter allem anderen im Dokument.
    schild_sections = [s for s in sections if s.find(cls="schild") is not None]
    ordered = [s for s in sections if s not in schild_sections] + schild_sections

    for section in ordered:
        conv.render_section(section)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    return conv


def main(argv):
    if len(argv) != 3:
        print(__doc__.strip())
        print("\nFehler: genau zwei Argumente erwartet "
              "(Quell-HTML und Ziel-DOCX).", file=sys.stderr)
        return 2
    src, dst = argv[1], argv[2]
    if not os.path.isfile(src):
        print("Quelldatei nicht gefunden: %s" % src, file=sys.stderr)
        return 1
    conv = convert(src, dst)
    print("geschrieben: %s" % dst)
    print("Bilder eingebettet: %d, uebersprungen: %d"
          % (conv.images, conv.skipped_images))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
